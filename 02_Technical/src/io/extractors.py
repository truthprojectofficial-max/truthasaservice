"""
Text extractors for .txt, .docx, and .pdf.  No cloud AI.
"""
import re
from pathlib import Path


def _clean_text(raw: str) -> str:
    if not raw:
        return ""
    raw = raw.replace("\x00", "")
    raw = raw.replace("\r\n", "\n").replace("\r", "\n")
    raw = re.sub(r"\n{3,}", "\n\n", raw)
    return raw.strip()


def extract_txt(path: Path) -> str:
    try:
        return _clean_text(path.read_text(encoding="utf-8"))
    except UnicodeDecodeError:
        return _clean_text(path.read_text(encoding="latin-1"))


def extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells if cell.text.strip())
    return _clean_text("\n".join(parts))


def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
    return _clean_text("\n".join(parts))


SUPPORTED_INPUTS = {
    ".txt": extract_txt,
    ".docx": extract_docx,
    ".pdf": extract_pdf,
}


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    extractor = SUPPORTED_INPUTS.get(ext)
    if extractor is None:
        raise ValueError(f"Unsupported input format: {ext} ({path})")
    return extractor(path)


def is_supported_input(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_INPUTS
